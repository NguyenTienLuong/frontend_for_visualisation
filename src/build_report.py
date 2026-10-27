"""Build the Vietnamese Cell2Cell management report from saved analysis outputs.

Run with the Codex bundled Python after the operation marker has been recorded.
This module only authors DOCX; Word/Poppler rendering and every-page visual QA are
performed separately.  No result in the report is taken from a remembered value.
"""
from __future__ import annotations

import argparse
import json
import math
import re
import sys
from collections import defaultdict
from datetime import date
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
ANALYSIS = ROOT / "output" / "analysis"
OUT = ROOT / "Cell2Cell_Customer_Churn_Report.docx"
ORANGE = "CC6B25"
BLACK = "000000"
GREY = "D9D9D9"


def number(value, decimals=0):
    if value is None or (isinstance(value, float) and not math.isfinite(value)):
        return "không xác định"
    return f"{float(value):,.{decimals}f}".replace(",", "~").replace(".", ",").replace("~", ".")


def percent(value, decimals=1):
    return number(float(value) * 100, decimals) + "%" if value is not None else "không xác định"


def load_json(path, default=None):
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else default


def field(paragraph, instruction, display="Cập nhật trường trong Word"):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " " + instruction + " "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = display
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)
    return run


def set_repeat_header(row):
    flag = OxmlElement("w:tblHeader")
    row._tr.get_or_add_trPr().append(flag)


class Report:
    def __init__(self, analysis=ANALYSIS, output=OUT):
        self.analysis = Path(analysis)
        self.output = Path(output)
        self.eda = load_json(self.analysis / "eda.json", {})
        self.summary = self.eda.get("summary", {})
        self.model = load_json(self.analysis / "model_summary.json", {})
        self.data_manifest = load_json(self.analysis / "data_manifest.json", {})
        self.risk_summary = load_json(self.analysis / "risk_summary.json", {})
        score_path = ROOT / "Cell2Cell_Churn_Risk_Scores.csv"
        self.candidate_stats = {}
        if score_path.exists():
            status = pd.read_csv(score_path, usecols=lambda c: c in ['priority_level', 'outreach_eligible_historical_status', 'observed_churn'])
            eligible = (status['outreach_eligible_historical_status'].astype(str).str.lower().eq('true')
                        if 'outreach_eligible_historical_status' in status else status['observed_churn'].eq(0))
            p1 = status['priority_level'].eq('P1')
            self.candidate_stats = {'all_p1': int(p1.sum()), 'candidate_p1': int((p1 & eligible).sum()), 'candidate_total': int(eligible.sum())}
        self.charts = load_json(self.analysis / "charts.json", [])
        model_charts = load_json(self.analysis / "model_charts.json", [])
        if isinstance(self.charts, dict):
            self.charts = self.charts.get("charts", list(self.charts.values()))
        if isinstance(model_charts, dict):
            model_charts = model_charts.get("charts", list(model_charts.values()))
        self.charts += model_charts
        self.used_charts = set()
        self.figure_counts = defaultdict(int)
        self.table_counts = defaultdict(int)
        self.figures = []
        self.tables = []
        self.doc = Document()
        self._styles()

    def _styles(self):
        section = self.doc.sections[0]
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin = section.bottom_margin = Cm(2)
        section.left_margin = section.right_margin = Cm(2)
        section.header_distance = section.footer_distance = Cm(0.8)
        styles = self.doc.styles
        normal = styles["Normal"]
        normal.font.name = "Arial"
        normal.font.size = Pt(10.5)
        normal.font.color.rgb = RGBColor.from_string(BLACK)
        normal.paragraph_format.line_spacing = 1.12
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.widow_control = True
        for name, size in [("Title", 23), ("Subtitle", 13), ("Heading 1", 17), ("Heading 2", 12.5), ("Heading 3", 11), ("Caption", 9)]:
            style = styles[name]
            style.font.name = "Arial"
            style.font.size = Pt(size)
            style.font.color.rgb = RGBColor.from_string(BLACK)
            style.font.bold = name not in ("Subtitle", "Caption")
            style.paragraph_format.space_before = Pt(9 if name.startswith("Heading") else 0)
            style.paragraph_format.space_after = Pt(6)
            style.paragraph_format.keep_with_next = name.startswith("Heading") or name == "Caption"
            pr = style.element.get_or_add_pPr()
            borders = pr.find(qn("w:pBdr"))
            if borders is not None:
                pr.remove(borders)
        styles["Caption"].paragraph_format.keep_with_next = False
        for style_name in ("TOC 1", "TOC 2"):
            if style_name in styles:
                styles[style_name].font.name = "Arial"
                styles[style_name].font.size = Pt(10)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer.add_run("Cell2Cell  |  ").font.size = Pt(8)
        field(footer, "PAGE", "1").font.size = Pt(8)
        settings = self.doc.settings.element
        update = OxmlElement("w:updateFields")
        update.set(qn("w:val"), "true")
        settings.append(update)
        self.doc.core_properties.title = "Phân tích khách hàng có khả năng rời bỏ dịch vụ và đề xuất giải pháp giữ chân khách hàng cho công ty viễn thông"
        self.doc.core_properties.subject = "Phân tích Customer Churn từ dữ liệu Cell2Cell"
        self.doc.core_properties.author = ""
        self.doc.core_properties.keywords = "Cell2Cell, churn, retention, customer analytics"

    def p(self, text, bold=False, size=None):
        p = self.doc.add_paragraph()
        r = p.add_run(str(text))
        r.bold = bold
        if size:
            r.font.size = Pt(size)
        return p

    def h(self, text, level=2):
        # Headings deliberately use words, digits and spaces only.
        text = re.sub(r"[.:—–/×&()]", " ", text)
        text = re.sub(r"\s+", " ", text).strip()
        return self.doc.add_heading(text, level)

    def reference(self, text, link_label, url):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(text + " ")
        run.font.size = Pt(9)
        link = OxmlElement('w:hyperlink')
        link.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
        r = OxmlElement('w:r')
        props = OxmlElement('w:rPr')
        size = OxmlElement('w:sz')
        size.set(qn('w:val'), '18')
        props.append(size)
        underline = OxmlElement('w:u')
        underline.set(qn('w:val'), 'single')
        props.append(underline)
        r.append(props)
        node = OxmlElement('w:t')
        node.text = link_label
        r.append(node)
        link.append(r)
        p._p.append(link)

    def page(self, heading=None, chapter=False):
        self.doc.add_page_break()
        if heading:
            self.h(heading, 1 if chapter else 2)

    def table(self, chapter, title, headers, rows, widths=None, font=9):
        self.table_counts[chapter] += 1
        label = f"Bảng {chapter}.{self.table_counts[chapter]} {title}"
        caption = self.doc.add_paragraph(label, "Caption")
        field(caption, f'TC "{label}" \\f T \\l 1', "")
        self.tables.append(label)
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.autofit = False
        if widths:
            for column, width in zip(table.columns, widths):
                column.width = Cm(width)
        for i, name in enumerate(headers):
            table.rows[0].cells[i].text = str(name)
        set_repeat_header(table.rows[0])
        for values in rows:
            cells = table.add_row().cells
            for i, value in enumerate(values):
                cells[i].text = str(value)
        for row_index, row in enumerate(table.rows):
            cant_split = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(cant_split)
            for i, cell in enumerate(row.cells):
                if widths:
                    cell.width = Cm(widths[i])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr = cell._tc.get_or_add_tcPr()
                borders = OxmlElement("w:tcBorders")
                for edge in ("top", "left", "bottom", "right"):
                    border = OxmlElement("w:" + edge)
                    border.set(qn("w:val"), "single")
                    border.set(qn("w:sz"), "4")
                    border.set(qn("w:color"), GREY)
                    borders.append(border)
                tcpr.append(borders)
                margins = OxmlElement("w:tcMar")
                for edge, val in (("top", "70"), ("bottom", "70"), ("left", "85"), ("right", "85")):
                    margin = OxmlElement("w:" + edge)
                    margin.set(qn("w:w"), val)
                    margin.set(qn("w:type"), "dxa")
                    margins.append(margin)
                tcpr.append(margins)
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "404040" if row_index == 0 else ("F4F4F4" if row_index % 2 == 0 else "FFFFFF"))
                tcpr.append(shading)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.line_spacing = 1.03
                    for run in p.runs:
                        run.font.name = "Arial"
                        run.font.size = Pt(font)
                        run.bold = row_index == 0
                        run.font.color.rgb = RGBColor.from_string("FFFFFF" if row_index == 0 else BLACK)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(1)
        return table

    def chart(self, chapter, *keys, width=16.8, max_height=9.3, interpretation=True):
        def normalized(value):
            return re.sub(r"[^a-z0-9]", "", str(value).lower())
        selected = None
        for key in keys:
            query = normalized(key)
            candidates = [c for c in self.charts if c.get("id") not in self.used_charts and query in normalized(c.get("id", "") + " " + c.get("slug", ""))]
            if not candidates:
                candidates = [c for c in self.charts if c.get("id") not in self.used_charts and query in normalized(c.get("title", "") + " " + c.get("section", ""))]
            if candidates:
                selected = next((c for c in candidates if c.get("report")), candidates[0])
                break
        if selected is None:
            raise ValueError(f"Required report chart unavailable: {keys}")
        path = Path(selected["path"])
        if not path.is_absolute():
            candidates = [ROOT / path, self.analysis / path, self.analysis / "charts" / path.name, ROOT / "output" / "figures" / path.name]
            path = next((p for p in candidates if p.exists()), candidates[0])
        if not path.exists():
            raise FileNotFoundError(path)
        self.used_charts.add(selected.get("id"))
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
        picture = p.add_run().add_picture(str(path), width=Cm(width))
        if picture.height > Cm(max_height):
            ratio = Cm(max_height) / picture.height
            picture.height = Cm(max_height)
            picture.width = int(picture.width * ratio)
        # Meaningful alt text travels with the final document.
        picture._inline.docPr.set("descr", selected.get("insight", selected.get("title", "")))
        self.figure_counts[chapter] += 1
        label = f"Hình {chapter}.{self.figure_counts[chapter]} {selected.get('title', selected.get('id', ''))}"
        caption = self.doc.add_paragraph(label, "Caption")
        field(caption, f'TC "{label.replace(chr(34), chr(39))}" \\f F \\l 1', "")
        self.figures.append(label)
        if interpretation:
            if selected.get("insight"):
                self.p(selected["insight"])
            if selected.get("business_meaning"):
                self.p(selected["business_meaning"])
        return selected

    def get_analysis(self, variable):
        return self.eda.get("analyses", {}).get(variable, {})

    def analysis_text(self, variable, question=None, include_business=True):
        a = self.get_analysis(variable)
        if a.get("key_insight"):
            self.p(f"{a.get('label', variable)}: {a['key_insight']}")
        if include_business and a.get("business_meaning"):
            self.p(a["business_meaning"])
        if question:
            self.p(question)

    def csv(self, name):
        path = self.analysis / name
        if not path.exists():
            path = ROOT / "output" / name
        return pd.read_csv(path) if path.exists() else pd.DataFrame()

    def row_summary(self, variable, max_rows=4):
        table = self.get_analysis(variable).get("table", [])
        if isinstance(table, dict):
            table = list(table.values())
        if not table:
            return ""
        lines = []
        for row in table[:max_rows]:
            label = row.get("group", row.get("label", row.get(variable, "")))
            n = row.get("n", row.get("count"))
            rate = row.get("churn_rate")
            if n is not None and rate is not None:
                lines.append(f"nhóm {label} có {number(n)} khách hàng và tỷ lệ churn {percent(rate)}")
        return "; ".join(lines)

    def intro_pages(self):
        p = self.doc.add_paragraph(style="Title")
        p.paragraph_format.space_before = Pt(68)
        p.add_run("PHÂN TÍCH KHÁCH HÀNG CÓ KHẢ NĂNG RỜI BỎ DỊCH VỤ VÀ ĐỀ XUẤT GIẢI PHÁP GIỮ CHÂN KHÁCH HÀNG CHO CÔNG TY VIỄN THÔNG")
        self.doc.add_paragraph("Phân tích dữ liệu Cell2Cell Customer Churn", "Subtitle")
        self.p("Báo cáo phân tích phục vụ lựa chọn khách hàng ưu tiên và thiết kế chương trình giữ chân", size=12)
        self.p("Ngày hoàn thành 27 tháng 10 năm 2026", size=10)
        self.p("Báo cáo dành cho ban quản lý và bộ phận Customer Retention. Phân tích kết hợp hồ sơ khách hàng, hành vi sử dụng, mô hình cảnh báo sớm và nguyên tắc đo lường hiệu quả can thiệp.")
        self.page("MỤC LỤC")
        field(self.doc.add_paragraph(), 'TOC \\o "1-1" \\h \\z \\u', "Mục lục tự động")
        self.page("DANH MỤC HÌNH")
        field(self.doc.add_paragraph(), 'TOC \\f F \\h \\z', "Danh mục hình tự động")
        self.page("DANH MỤC BẢNG")
        field(self.doc.add_paragraph(), 'TOC \\f T \\h \\z', "Danh mục bảng tự động")

    def executive(self):
        s, m = self.summary, self.model
        test = m["test"]
        targeting = self.csv("targeting_test.csv")
        top10 = targeting.iloc[(targeting["capacity"] - 0.1).abs().argsort()[:1]].iloc[0] if not targeting.empty and "capacity" in targeting else None
        top20 = targeting.iloc[(targeting["capacity"] - 0.2).abs().argsort()[:1]].iloc[0] if not targeting.empty and "capacity" in targeting else None
        self.page("TÓM TẮT ĐIỀU HÀNH", True)
        self.p(f"Dữ liệu khách hàng của công ty ghi nhận {number(s['n_customers'])} hồ sơ, trong đó {number(s['n_churn'])} khách hàng có nhãn rời bỏ dịch vụ, tương ứng {percent(s['churn_rate'])}. Đây là tỷ lệ quan sát trong ảnh chụp dữ liệu lịch sử đang phân tích. Con số này mô tả tập khách hàng có nhãn và không được diễn giải thành tỷ lệ churn của toàn bộ thuê bao đang hoạt động trong một tháng hiện tại.")
        leaders = [c for c in self.charts if c.get("insight") and c.get("slug") in ["MonthsInService", "PercChangeMinutes", "CurrentEquipmentDays"]][:3]
        for c in leaders:
            label = self.get_analysis(c.get('slug', '')).get('label', c.get('slug', 'Tín hiệu khách hàng'))
            self.p(f"{label}: {c['insight']}")
        self.p(f"Mô hình được lựa chọn là {m['selected_model']}. Trên tập kiểm tra độc lập, ROC AUC đạt {number(test['roc_auc'], 3)}, Average Precision đạt {number(test['pr_auc_average_precision'], 3)}; tại ngưỡng chọn bằng tập validation, precision là {percent(test['precision'])} và recall là {percent(test['recall'])}. Kết quả này đánh giá khả năng sắp hạng rủi ro trên dữ liệu lịch sử, chưa đo tác động của một ưu đãi giữ chân.")
        if top10 is not None:
            self.p(f"Nếu liên hệ nhóm {percent(top10['capacity'], 0)} khách hàng có điểm cao nhất trên tập kiểm tra, precision đạt {percent(top10['precision'])} và lift đạt {number(top10['lift'], 2)} lần mức chọn ngẫu nhiên. " + (f"Mở rộng tới {percent(top20['capacity'], 0)} thu được {percent(top20['capture_recall'])} tổng số khách hàng churn trong tập kiểm tra." if top20 is not None else ""))
        self.p("Công ty nên ưu tiên khách hàng đồng thời có rủi ro cao và giá trị doanh thu tháng cao; tiếp theo là khách hàng có dấu hiệu trải nghiệm cần được xử lý bằng một hành động cụ thể. Doanh thu tháng được dùng như đại diện cho giá trị khách hàng, không phải lợi nhuận hay giá trị vòng đời. Với khách hàng giá trị thấp, nên bắt đầu bằng kênh tự phục vụ hoặc liên hệ có chi phí thấp và giữ nguyên nhóm đối chứng để đo hiệu quả tăng thêm.")
        if self.candidate_stats:
            c = self.candidate_stats
            self.p(f"Nhóm P1 gồm {number(c['all_p1'])} hồ sơ lịch sử có giá trị cao và rủi ro cao. Sau khi loại hồ sơ đã có nhãn churn, còn {number(c['candidate_p1'])} ứng viên P1 trong {number(c['candidate_total'])} hồ sơ chưa churn theo trạng thái lịch sử. Đây vẫn là ứng viên cần xác minh trạng thái tại ngày triển khai, không phải danh sách gọi ngay.")
        self.p("Các biến phát sinh trong quá trình retention và các mã định danh được loại khỏi mô hình cảnh báo sớm. Danh sách điểm số được tạo từ dự đoán ngoài mẫu của ảnh chụp lịch sử; trước khi sử dụng vận hành cần chấm lại tập thuê bao còn hoạt động bằng dữ liệu sẵn có trước thời điểm liên hệ. Mỗi chiến dịch cần ghi nhận phân nhóm thử nghiệm, tiếp xúc, ưu đãi, chi phí và trạng thái thuê bao sau can thiệp.")
        self.p("Quyết định đề nghị là triển khai thử nghiệm có đối chứng theo phân khúc rủi ro và giá trị. Chỉ số ưu tiên gồm churn rate, precision trong nhóm được liên hệ, số churn được bao phủ, incremental retention lift, chi phí cho mỗi khách hàng giữ lại thêm và doanh thu được bảo vệ. Ngân sách chỉ được phê duyệt sau khi có biên lợi nhuận, chi phí liên hệ và tác động thực nghiệm của từng chính sách.")

    def chapter1(self):
        self.page("CHƯƠNG 1 ĐẶT VẤN ĐỀ", True)
        self.p("Khách hàng rời bỏ làm giảm nguồn doanh thu lặp lại và buộc công ty dành thêm nguồn lực để thay thế thuê bao đã mất. Quyết định giữ chân vì vậy cần bắt đầu từ một câu hỏi vận hành cụ thể: với năng lực liên hệ hữu hạn, bộ phận Customer Retention nên xem xét khách hàng nào trước và có căn cứ gì để chọn hành động cho họ.")
        self.h("1 1 Bài toán kinh doanh và phạm vi quyết định")
        self.p("Churn trong báo cáo là trạng thái rời bỏ được ghi trong biến mục tiêu của dữ liệu Cell2Cell. Phân tích tìm các đặc điểm liên quan đến trạng thái này và xây dựng một thứ tự ưu tiên. Dữ liệu không ghi lại đầy đủ thời điểm quan sát và toàn bộ lịch sử tiếp xúc, vì vậy báo cáo không gán một khoảng dự báo tương lai cụ thể và không tự tạo tỷ lệ churn theo tháng.")
        self.p("Giữ chân hiệu quả đòi hỏi phân biệt rủi ro rời bỏ, giá trị kinh tế và khả năng phản ứng với can thiệp. Một khách hàng có điểm rủi ro cao chưa chắc cần ưu đãi lớn; họ có thể cần giải quyết lỗi mạng hoặc hỗ trợ sử dụng. Ngược lại, một khách hàng doanh thu cao chỉ nên được liên hệ khi lợi ích tăng thêm có thể bù chi phí của chính sách.")
        self.h("1 2 Mục tiêu và câu hỏi phân tích")
        self.p("Dự án đo quy mô churn trong tập quan sát; mô tả khách hàng theo vòng đời, doanh thu, mức sử dụng, chất lượng cuộc gọi, thiết bị và hồ sơ; xác định các tổ hợp tín hiệu; đánh giá mô hình cảnh báo sớm; và chuyển kết quả thành quy tắc phân bổ nguồn lực. Mỗi kết luận mô tả phải đi kèm quy mô nhóm và tỷ lệ churn để tránh ưu tiên một nhóm chỉ vì nhóm đó đông khách hàng.")
        self.p("Các câu hỏi chính là churn tập trung ở giai đoạn quan hệ nào; liệu biến động sử dụng và doanh thu có cung cấp cảnh báo; các chỉ báo cuộc gọi và hỗ trợ có còn hữu ích khi kết hợp với vòng đời; và mô hình bao phủ được bao nhiêu khách hàng churn khi chỉ lựa chọn một phần tập khách hàng. Phần chính sách đặt thêm câu hỏi liệu một hành động có tạo ra giữ chân tăng thêm sau khi tính chi phí.")
        self.h("1 3 Quy trình và tiêu chuẩn bằng chứng")
        self.p("Quy trình bắt đầu bằng nhận diện cấu trúc và kiểm tra chất lượng dữ liệu, sau đó phân tích so sánh, phân tích kết hợp và xếp hạng tín hiệu. Mô hình được huấn luyện sau khi loại biến có nguy cơ rò rỉ; hiệu quả chọn khách hàng được kiểm tra trên phần dữ liệu chưa dùng để chọn mô hình. Điểm rủi ro được ghép với đại diện giá trị để thiết kế phân khúc, chính sách và KPI.")
        self.p("Các biểu đồ dùng vị trí và chiều dài để so sánh, màu cam để tập trung vào churn và nền xám cho nhóm duy trì dịch vụ. Mỗi hình trả lời một câu hỏi cụ thể; chú giải, nhãn và chú thích được giới hạn ở thông tin cần cho quyết định. Tương tác trong notebook hỗ trợ nhận diện, lọc và so sánh từng khách hàng, còn báo cáo chọn các hình cần thiết để truyền đạt kết luận.")
        self.p("Kết luận chương 1. Mục tiêu của dự án là hỗ trợ một quyết định liên hệ có giới hạn nguồn lực. Bằng chứng quan sát được dùng để đặt ưu tiên và xây dựng giả thuyết; tác động thực sự của chính sách phải được kiểm chứng bằng thử nghiệm.")

    def chapter2(self):
        s = self.summary
        self.page("CHƯƠNG 2 TỔNG QUAN DỮ LIỆU VÀ PHƯƠNG PHÁP", True)
        self.p("Dữ liệu được đọc từ nguồn Cell2Cell qua bộ tải empulse theo hướng dẫn đã cung cấp. Cấu trúc và kiểu dữ liệu được nhận diện từ dữ liệu tải về, thay vì áp một danh sách cột cố định từ phiên bản khác. Các bảng và hình trong báo cáo cùng sử dụng một bộ kết quả phân tích đã lưu để bảo đảm đồng nhất với notebook.")
        manifest = self.data_manifest
        if manifest:
            self.p(f"Nguồn gốc có {number(manifest['raw_rows'])} dòng và {number(manifest['raw_columns'])} cột. Bộ tải nguồn loại {number(manifest['rows_removed_missing_monthly_revenue'])} dòng thiếu MonthlyRevenue, tạo cohort phân tích {number(manifest['loader_rows'])} khách hàng. Toàn bộ tỷ lệ trong EDA và mô hình của báo cáo dùng cohort này, trừ phần kiểm tra dữ liệu nguồn được ghi rõ. Quy tắc loại dòng là xử lý của nguồn và có thể ảnh hưởng tính đại diện.")
            if manifest.get('rows_removed_missing_monthly_revenue'):
                self.p(f"Trong phần bị loại, {number(manifest['removed_churn_count'])} trên {number(manifest['rows_removed_missing_monthly_revenue'])} khách hàng có nhãn churn, tương ứng {percent(manifest['removed_churn_count'] / manifest['rows_removed_missing_monthly_revenue'])}, so với {percent(s['churn_rate'])} ở cohort còn lại. Vì vậy không suy rộng trực tiếp tỷ lệ cohort cho toàn bộ nguồn hoặc cơ sở thuê bao hiện tại.")
        self.table(2, "Quy mô và cấu trúc tập quan sát", ["Chỉ tiêu", "Giá trị", "Diễn giải"], [
            ["Hồ sơ khách hàng", number(s['n_customers']), "Tập có nhãn sử dụng cho phân tích"],
            ["Biến đầu vào", number(s['n_features']), "Không bao gồm nhãn mục tiêu"],
            ["Biến số", number(s['n_numeric']), "Theo kiểu dữ liệu đã nhận diện"],
            ["Biến phân loại", number(s['n_categorical']), "Bao gồm các trường mã hóa phân nhóm"],
            ["Khách hàng churn", number(s['n_churn']), "Số hồ sơ có nhãn mục tiêu bằng một"],
            ["Tỷ lệ churn", percent(s['churn_rate']), "Mẫu số là toàn bộ tập có nhãn"],
        ], [5.0, 3.0, 8.8])
        self.h("2 1 Đơn vị quan sát và ý nghĩa biến")
        self.p("Đơn vị quan sát là một hồ sơ khách hàng trong ảnh chụp dữ liệu. Các chỉ báo mức sử dụng, doanh thu và chất lượng cuộc gọi được hiểu theo nhãn biến và mô tả đi kèm nguồn dữ liệu. Các trường thay đổi như PercChangeMinutes và PercChangeRevenues được giữ theo đơn vị do nguồn cung cấp; tên có chữ PercChange không đủ để tự chuyển đổi mọi giá trị thành phần trăm.")
        variables = self.eda.get("variables", {})
        important = ["MonthlyRevenue", "MonthlyMinutes", "MonthsInService", "TotalRecurringCharge", "PercChangeMinutes", "PercChangeRevenues", "DroppedCalls", "BlockedCalls", "CustomerCareCalls", "CurrentEquipmentDays"]
        meanings = {
            "MonthlyRevenue": "Doanh thu tháng theo trường nguồn", "MonthlyMinutes": "Mức sử dụng phút gọi", "MonthsInService": "Thời gian quan hệ tính bằng tháng", "TotalRecurringCharge": "Mức phí định kỳ", "PercChangeMinutes": "Thay đổi sử dụng theo định nghĩa nguồn", "PercChangeRevenues": "Thay đổi doanh thu theo định nghĩa nguồn", "DroppedCalls": "Chỉ báo cuộc gọi bị rớt", "BlockedCalls": "Chỉ báo cuộc gọi bị chặn", "CustomerCareCalls": "Mức liên hệ chăm sóc khách hàng", "CurrentEquipmentDays": "Số ngày sử dụng thiết bị hiện tại"}
        rows = []
        for name in important:
            if name in variables or name in self.eda.get("analyses", {}):
                val = variables.get(name, {})
                group = val.get("group", "Biến vận hành") if isinstance(val, dict) else val
                group = {'Revenue and charges':'Doanh thu và phí', 'Usage and calling':'Sử dụng và cuộc gọi', 'Customer relationship':'Quan hệ với công ty', 'Service experience':'Trải nghiệm dịch vụ', 'Devices':'Thiết bị', 'Demographics':'Hồ sơ khách hàng', 'Credit profile':'Hồ sơ tín dụng'}.get(group,group)
                rows.append([name, group, meanings[name]])
        self.table(2, "Từ điển các biến chính", ["Biến", "Nhóm", "Ý nghĩa phân tích"], rows, [5.2, 4.2, 7.4], 8.5)
        self.page("2 2 Phương pháp và cách đọc bằng chứng")
        self.p("Phân tích mô tả so sánh các nhóm theo quy mô, số churn và tỷ lệ churn. Với biến liên tục, các khoảng được xây từ phân phối quan sát hoặc một ngưỡng có nghĩa kinh doanh được ghi rõ. Nhóm thiếu dữ liệu được giữ riêng trong phân tích mô tả khi cần, giúp tránh bỏ qua các khách hàng không có giá trị đo lường.")
        self.table(2, "Phương pháp gắn với quyết định", ["Phương pháp", "Câu hỏi", "Giới hạn cần giữ"], [
            ["Thống kê và phân phối", "Đặc điểm điển hình và mức phân tán", "Trung bình có thể chịu ảnh hưởng giá trị cực trị"],
            ["Tỷ lệ churn theo nhóm", "Rủi ro quan sát tập trung ở đâu", "Luôn đọc cùng quy mô nhóm và khoảng bất định"],
            ["Phân tích đa biến", "Tổ hợp tín hiệu có khác biệt gì", "Quan hệ quan sát không xác nhận nguyên nhân"],
            ["Đánh giá mô hình", "Sắp hạng và lựa chọn được ai", "Tập kiểm tra lịch sử không bảo đảm tương lai"],
            ["Rủi ro và giá trị", "Ưu tiên nguồn lực cho ai", "Doanh thu tháng chưa phải biên lợi nhuận"],
            ["Thử nghiệm đối chứng", "Chính sách tạo thêm giữ chân hay không", "Cần ngẫu nhiên hóa và theo dõi kết quả"],
        ], [4, 5.8, 7])
        self.h("2 3 Nguyên tắc trực quan hóa")
        self.p("Tư duy Storytelling with Data được áp dụng bằng cách bắt đầu từ người ra quyết định và câu hỏi kinh doanh, chọn hình phù hợp, giảm chi tiết trang trí và dùng màu có chủ đích. Các thanh tỷ lệ đặt cạnh nhãn nhóm giúp so sánh trực tiếp; biểu đồ phân phối giữ thông tin về độ lệch; đường biểu diễn chỉ dùng cho thứ tự có nghĩa; heatmap dành cho tương tác cần đọc theo hai chiều.")
        self.p("Theo tư duy Interactive Data Visualization, bốn công cụ khám phá trong notebook có nhiệm vụ rõ ràng: xác định khách hàng có điểm và giá trị đáng chú ý; so sánh sử dụng với doanh thu; lọc các thay đổi hành vi; và kiểm tra phân phối điểm. Hover hiển thị thuộc tính cần thiết để nhận diện, còn lọc và chọn giúp chuyển từ tổng quan sang trường hợp cụ thể. Dữ liệu tương tác vẫn là dữ liệu lịch sử trong phạm vi phân tích.")
        self.h("2 4 Nguồn và khả năng tái lập")
        self.p("Nguồn dữ liệu chính là bộ Cell2Cell được bộ tải empulse cung cấp. Nguồn phương pháp trực quan là Cole Nussbaumer Knaflic, Storytelling with Data, và Matthew O Ward, Georges Grinstein, Daniel Keim, Interactive Data Visualization Foundations Techniques and Applications. Các tài liệu định hướng cách biểu diễn và thao tác khám phá; mọi số đo khách hàng trong báo cáo được tính từ dữ liệu phân tích của công ty.")
        self.p("Kết luận chương 2. Cấu trúc thực tế của dữ liệu quyết định phạm vi phân tích. Báo cáo thống nhất đơn vị quan sát, mẫu số và giới hạn diễn giải trước khi xem xét các khác biệt giữa churn và nhóm duy trì dịch vụ.")

    def chapter3(self):
        q = self.eda.get("quality", {})
        self.page("CHƯƠNG 3 KIỂM TRA VÀ TIỀN XỬ LÝ DỮ LIỆU", True)
        self.p("Kiểm tra chất lượng nhằm xác định một giá trị có thể sử dụng cho phân tích và mô hình hay không. Thiếu dữ liệu, mã phân loại và giá trị cực trị được xử lý theo ý nghĩa của từng trường; không xóa khách hàng chỉ vì mức sử dụng hoặc doanh thu cao hơn phần lớn tập quan sát.")
        self.chart(3, "E02")
        if "analysis_duplicate_rows" in q:
            self.p(f"Cohort có {number(q['analysis_duplicate_rows'])} dòng trùng toàn bộ và {number(q.get('duplicate_customer_ids', 0))} mã khách hàng bị lặp; tập nguồn có {number(q.get('raw_duplicate_rows', 0))} dòng trùng toàn bộ. Hai phép kiểm tra được tách riêng vì một mã lặp có thể mang nhiều bản ghi khác nhau.")
        dup = q.get("duplicates", {})
        if isinstance(dup, dict):
            rows_dup = dup.get("duplicate_rows", dup.get("rows", dup.get("full_rows")))
            if rows_dup is not None:
                self.p(f"Kiểm tra trùng toàn bộ dòng ghi nhận {number(rows_dup)} hồ sơ trùng. Kiểm tra định danh được thực hiện riêng khi nguồn có mã khách hàng; các chỉ số liên quan được lưu trong bảng chất lượng dữ liệu của notebook.")
        self.h("3 1 Giá trị thiếu và cực trị")
        self.p("Phân tích mô tả không che giấu phần thiếu. Trong mô hình, giá trị số được điền bằng trung vị học từ tập huấn luyện và kèm chỉ báo thiếu khi bộ xử lý hỗ trợ; biến phân loại được gán một trạng thái thiếu riêng. Cách làm này tránh đưa thông tin từ tập kiểm tra vào tham số điền khuyết và giữ được khả năng thiếu dữ liệu cũng là một tín hiệu.")
        self.p("Các mã đặc biệt được xử lý theo trường: tuổi hộ gia đình bằng không được chuyển thành thiếu; IncomeGroup mã không thành Unknown và các mã còn lại được coi là phân loại; HandsetPrice được chuyển sang số với Unknown thành thiếu. Mọi bước điền trung vị, chỉ báo thiếu, chuẩn hóa và mã hóa đều chỉ học trong phần huấn luyện tương ứng.")
        self.p("Giá trị ngoài khoảng IQR là dấu hiệu cần rà soát, không phải bằng chứng lỗi. Doanh thu và lưu lượng cuộc gọi có thể lệch mạnh vì khách hàng dùng cường độ cao. Mô hình cây xử lý quan hệ phi tuyến trực tiếp; mô hình tuyến tính sử dụng bước chuẩn hóa học từ tập huấn luyện. Dữ liệu gốc được giữ để kiểm tra và giải thích.")
        negative_equipment = next((row['n_negative'] for row in q.get('validity', []) if row['variable'] == 'CurrentEquipmentDays'), None)
        if negative_equipment:
            self.p(f"CurrentEquipmentDays có {number(negative_equipment)} giá trị âm cần xác minh định nghĩa hoặc cách ghi nhận. Các hồ sơ này được giữ và đánh dấu để rà soát, không tự gán là lỗi rồi xóa. Doanh thu âm cũng được giữ vì nguồn chưa đủ căn cứ phân biệt khoản điều chỉnh tài khoản với lỗi.")
        self.page("3 2 Tính hợp lệ theo thời điểm và rò rỉ dữ liệu")
        leak = self.model.get("excluded_potential_leakage", self.model.get("excluded_leakage", []))
        ids = self.model.get("excluded_identifiers", [])
        self.p("Mô hình vận hành được định nghĩa tại thời điểm trước khi khách hàng được đưa vào xử lý retention. Một biến phản ánh việc đã gọi bộ phận giữ chân hoặc đã nhận ưu đãi có thể xuất hiện cùng quá trình phát hiện ý định rời bỏ. Không thể khẳng định biến này sẵn có trước quyết định can thiệp chỉ dựa vào tương quan với nhãn churn.")
        self.table(3, "Các quyết định xử lý phục vụ cảnh báo sớm", ["Vấn đề", "Quyết định", "Lý do"], [
            ["Biến retention", ", ".join(leak) if leak else "Loại các trường có nguy cơ phát sinh trong quá trình giữ chân", "Không chứng minh được tính sẵn có trước can thiệp"],
            ["Mã định danh", ", ".join(ids) if ids else "Loại mã định danh khỏi đặc trưng", "Dùng để đối chiếu hồ sơ thay vì suy luận rủi ro"],
            ["Giá trị số thiếu", "Trung vị từ tập huấn luyện", "Giảm ảnh hưởng đuôi phân phối và tránh rò rỉ tập kiểm tra"],
            ["Phân loại thiếu", "Trạng thái thiếu riêng", "Bảo toàn quy mô và khả năng nhận diện thiếu dữ liệu"],
            ["Giá trị cực trị", "Giữ lại và mô tả", "Phân biệt hành vi thật với lỗi nhập dữ liệu"],
            ["Ngưỡng và hiệu chỉnh", "Chọn trên validation", "Không tối ưu quyết định trên tập test"],
        ], [3.4, 7.0, 6.4], 8.5)
        self.p("Các trường retention vẫn được dùng để mô tả quy trình hỗ trợ. Tỷ lệ churn cao trong nhóm đã liên hệ retention có thể phản ánh việc công ty chọn đúng các khách hàng khó giữ, và không được diễn giải là liên hệ retention làm tăng churn. Đánh giá ưu đãi cần phân nhóm thử nghiệm hoặc thiết kế nhân quả phù hợp.")
        self.h("3 3 Phân tách dữ liệu và kiểm soát quy trình")
        self.p("Mọi bước có tham số học từ dữ liệu được đặt trong pipeline. Tham số điền thiếu, chuẩn hóa và mã hóa được fit trên tập huấn luyện; validation phục vụ chọn mô hình, ngưỡng và hiệu chỉnh; test chỉ phục vụ báo cáo cuối. Việc chia có phân tầng theo nhãn giúp giữ cấu trúc churn của từng phần nhưng không thay thế đánh giá theo thời gian khi triển khai.")
        self.p("Kết luận chương 3. Tập phân tích mô tả bảo toàn thông tin cần thiết, còn mô hình cảnh báo sớm chỉ dùng các đặc trưng được chấp nhận trong phạm vi thời điểm dự đoán. Hạn chế lớn còn lại là thiếu lịch sử dọc theo thời gian để xác minh đầy đủ thời điểm xuất hiện của từng trường.")

    def chapter4(self):
        s = self.summary
        self.page("CHƯƠNG 4 BỨC TRANH KHÁCH HÀNG VÀ CHURN", True)
        self.p(f"Trong {number(s['n_customers'])} hồ sơ được quan sát, số khách hàng churn là {number(s['n_churn'])}; nhóm còn lại có {number(s['n_customers'] - s['n_churn'])} hồ sơ. Tỷ lệ churn bằng số khách hàng có nhãn rời bỏ chia cho toàn bộ số hồ sơ có nhãn, đạt {percent(s['churn_rate'])}. Mức nền này được dùng để đánh giá lift và so sánh các phân khúc trong những chương sau.")
        self.chart(4, "E01", width=15.8)
        self.h("4 1 Vòng đời và nguồn doanh thu")
        self.analysis_text("MonthsInService")
        self.analysis_text("MonthlyRevenue")
        if s.get("monthly_revenue_total") is not None:
            self.p(f"Tổng MonthlyRevenue ghi nhận trong ảnh chụp là {number(s['monthly_revenue_total'], 2)} theo đơn vị tiền của nguồn dữ liệu. " + (f"Phần gắn với các hồ sơ churn là {number(s['churn_monthly_revenue_proxy'], 2)}. " if s.get("churn_monthly_revenue_proxy") is not None else "") + "Đây là quy mô doanh thu tháng gắn với hồ sơ quan sát, không phải doanh thu đã mất trong một khoảng thời gian hay giá trị có thể thu hồi chắc chắn.")
        self.p("Số churn lớn và tỷ lệ churn cao trả lời hai câu hỏi khác nhau. Nhóm đông khách hàng có thể tạo phần lớn số rời bỏ dù tỷ lệ chỉ gần mức nền; nhóm nhỏ có tỷ lệ cao cần được xem xét cùng quy mô và khả năng can thiệp trước khi phân bổ nhân sự.")
        self.page("4 2 Phân phối sử dụng và khác biệt giữa các khách hàng")
        self.chart(4, "E05", max_height=7, interpretation=False)
        self.chart(4, "E04", width=15.8, max_height=7, interpretation=False)
        self.p("Phân phối cho thấy vì sao một chỉ số bình quân không đủ để mô tả tập khách hàng. Với lưu lượng hoặc doanh thu có đuôi dài, một thay đổi ở nhóm sử dụng rất cao có thể tác động đến trung bình mà không đại diện cho phần lớn thuê bao. Phần phân tích chi tiết vì vậy đọc đồng thời phân phối, nhóm theo khoảng và tỷ lệ churn.")
        self.p("Kết luận chương 4. Bức tranh chung xác lập quy mô churn và mức đa dạng trong vòng đời, giá trị và sử dụng. Câu hỏi tiếp theo là những khác biệt nào đủ lớn, có quy mô phù hợp và có một hành động nghiệp vụ tương ứng.")

    def detail_page(self, heading, keys, variables, interpretation, action, closing, compact=False):
        self.page(heading)
        self.chart(5, *keys, interpretation=False)
        for var in variables:
            self.analysis_text(var, include_business=not compact)
        self.p(interpretation)
        self.p(action)
        self.p(closing)

    def chapter5(self):
        self.page("CHƯƠNG 5 CÁC ĐẶC TRƯNG LIÊN QUAN ĐẾN CHURN", True)
        self.p("Chương này tổ chức các biến theo vấn đề kinh doanh. Các khác biệt được mô tả như quan hệ quan sát và được dùng để xác định giả thuyết hành động. Một tỷ lệ cao chỉ trở thành ưu tiên khi đi kèm quy mô đáng kể, thông tin sẵn có và một biện pháp có thể kiểm chứng.")
        self.h("5 1 Thời gian gắn bó với công ty")
        self.chart(5, "E06", interpretation=False)
        self.analysis_text("MonthsInService")
        self.p("Cách đọc vòng đời cần phân biệt độ dài quan hệ hiện tại với nguy cơ tại từng thời điểm. Khách hàng lâu năm trong ảnh chụp đã có cơ hội duy trì quan hệ lâu hơn; sự khác biệt giữa các nhóm không phải xác suất sống sót theo tháng. Muốn kiểm tra nguy cơ theo hành trình cần bổ sung ngày kích hoạt, ngày rời bỏ và lịch sử trạng thái thuê bao.")
        self.p("Bộ phận retention nên dùng kết quả này để điều chỉnh nội dung chăm sóc theo giai đoạn quan hệ. Hỗ trợ kích hoạt và hiểu cước phù hợp với khách hàng mới; rà soát mức sử dụng, thiết bị và quyền lợi phù hợp với khách hàng đã gắn bó. Các khoảng phân nhóm được lấy từ dữ liệu và cần được cập nhật cùng thay đổi cơ cấu thuê bao.")
        self.p("Kết quả vòng đời là điểm khởi đầu cho phân tích kết hợp. Chương 6 kiểm tra nó cùng biến động sử dụng và thiết bị để xác định liệu một tín hiệu đơn lẻ có đang phản ánh đặc điểm của nhóm khác.")
        self.detail_page("5 2 Doanh thu tháng và mức phí định kỳ", ["E07"], ["MonthlyRevenue", "TotalRecurringCharge"],
            "Doanh thu tháng vừa liên quan tới cường độ sử dụng vừa phản ánh cấu trúc gói và các khoản phát sinh. Mức phí định kỳ giúp tách một phần cam kết nền khỏi chi tiêu vượt gói. So sánh hai trường này có giá trị để đặt câu hỏi về mức phù hợp của gói, nhưng chưa đủ để kết luận khách hàng không hài lòng với giá.",
            "Trước khi đề nghị giảm giá, công ty cần rà soát thành phần hóa đơn, mức sử dụng và biên lợi nhuận của từng hồ sơ. Khách hàng doanh thu cao cần một chính sách bảo vệ trải nghiệm và giá trị; một ưu đãi đại trà cho mọi khách hàng có mức phí cao có thể giảm doanh thu mà không tạo giữ chân tăng thêm.",
            "Cần theo dõi cả tỷ lệ churn và phần doanh thu gắn với từng nhóm. Mục tiêu tiếp theo là kiểm tra doanh thu cao do sử dụng thật hay do chi phí phát sinh chưa phù hợp.")
        self.detail_page("5 3 Mức sử dụng và cấu trúc cuộc gọi", ["E09"], ["MonthlyMinutes", "InboundCalls", "OutboundCalls", "ReceivedCalls"],
            "Mức sử dụng phản ánh vai trò của dịch vụ trong sinh hoạt hoặc công việc của khách hàng. Lưu lượng thấp có thể phù hợp với nhu cầu thật, phản ánh thuê bao phụ hoặc báo hiệu dịch vụ ít được dùng. Lưu lượng cao có thể gắn với giá trị lớn nhưng cũng có thể đi kèm bất tiện hoặc cước phát sinh. Những cơ chế này cần được kiểm tra bằng dữ liệu bổ sung.",
            "Nhóm phân tích nên xem xét cấu trúc cuộc gọi vào, ra và nhận được cùng tổng phút gọi khi tư vấn gói. Hành động phù hợp là rà soát mức tương thích của gói và cách sử dụng; không nên tăng khuyến mại chỉ vì lưu lượng thấp. Nhật ký phản hồi sau liên hệ giúp phân biệt thiếu nhu cầu với trải nghiệm chưa tốt.",
            "Phân phối sử dụng tĩnh được bổ sung bằng thay đổi gần đây ở trang tiếp theo. Mức hiện tại và hướng thay đổi trả lời hai câu hỏi khác nhau về tình trạng quan hệ.")
        self.page("5 4 Thay đổi sử dụng và doanh thu gần đây")
        self.chart(5, "E10", max_height=9.3, interpretation=False)
        self.analysis_text("PercChangeMinutes")
        self.analysis_text("PercChangeRevenues")
        self.p("Một mức sử dụng giảm có thể phản ánh thay đổi nhu cầu hoặc trải nghiệm, nhưng ảnh chụp không xác định nguyên nhân. Các trường thay đổi được giữ theo đơn vị nguồn, không tự chuyển thành phần trăm. Thiếu dữ liệu được giữ riêng vì có thể phản ánh cơ chế đo lường khác nhau. Khi liên hệ, cần xác minh nhu cầu, lịch sử gói và cước trước khi đề nghị ưu đãi.")
        self.detail_page("5 5 Vượt gói và sử dụng roaming", ["E12"], ["OverageMinutes", "RoamingCalls"],
            "Vượt gói và roaming có thể tạo khoản chi không đều giữa các khách hàng. Chúng cũng có thể đơn giản phản ánh nhu cầu cao hoặc di chuyển nhiều. Cần đọc số khách hàng trong mỗi nhóm để tránh diễn giải một nhóm hiếm như một vấn đề phổ biến trên toàn công ty.",
            "Hành động nên bắt đầu từ kiểm tra sự phù hợp của gói và độ rõ ràng của hóa đơn. Cảnh báo chi tiêu và giải thích lựa chọn gói có thể là phương án ít tốn kém hơn giảm giá. Khi thiết kế ưu đãi roaming, cần tính chi phí cung cấp dịch vụ và xem xét doanh thu đóng góp thay vì chỉ nhìn tổng doanh thu.",
            "Khác biệt churn theo nhóm là căn cứ chọn đối tượng cần chẩn đoán. Để chứng minh chính sách về gói cước hiệu quả, công ty cần một thử nghiệm có chi phí và tiêu chí giữ chân được xác định trước.")
        self.detail_page("5 6 Chất lượng cuộc gọi và liên hệ với churn", ["E14"], ["DroppedCalls", "BlockedCalls", "UnansweredCalls", "DroppedBlockedCalls"],
            "Số cuộc gọi rớt hoặc bị chặn chịu ảnh hưởng cả chất lượng mạng lẫn mức sử dụng. Khách hàng gọi nhiều có nhiều cơ hội gặp sự cố hơn; vì vậy so sánh số đếm thô chưa đủ để quy trách nhiệm cho mạng. Chỉ báo tổng hợp có thể hữu ích để sàng lọc nhưng không thay thế tỷ lệ lỗi trên số cuộc gọi và dữ liệu theo khu vực.",
            "Công ty nên kiểm tra hồ sơ sự cố và vùng phủ trước khi ưu đãi tài chính cho khách hàng có tín hiệu chất lượng. Giải quyết nguyên nhân kỹ thuật cần được ưu tiên khi xác nhận có lỗi; bộ phận chăm sóc phải thông báo tiến độ và ghi nhận vấn đề đã được khắc phục. Kết quả liên hệ cần được nối với dữ liệu mạng để đánh giá.",
            "Các quan hệ quan sát không chứng minh rằng một cuộc gọi rớt làm khách hàng rời bỏ. Chúng xác định nhóm nên được điều tra và đặt nền cho phân tích kết hợp chất lượng với mức sử dụng.")
        self.detail_page("5 7 Chăm sóc khách hàng và quy trình retention", ["E17"], ["CustomerCareCalls", "RetentionCalls", "RetentionOffersAccepted", "MadeCallToRetentionTeam"],
            "Liên hệ chăm sóc có thể đại diện cho một yêu cầu chưa được giải quyết, nhưng cũng có thể là tương tác bình thường. Số cuộc gọi không cho biết nội dung, mức nghiêm trọng hoặc kết quả xử lý. Tương tự, trạng thái đã được retention liên hệ chịu ảnh hưởng bởi quy tắc lựa chọn khách hàng của công ty, tạo ra khác biệt giữa nhóm được và chưa được chăm sóc.",
            "Nhóm liên hệ hỗ trợ nhiều nên được rà soát lịch sử và giải quyết tận gốc trước khi gửi thêm ưu đãi. Hệ thống cần ghi lý do liên hệ, số lần lặp lại, thời gian xử lý, lần chuyển bộ phận và xác nhận giải quyết. Các biến retention chỉ phục vụ mô tả quy trình trong phần này và đã được loại khỏi mô hình cảnh báo sớm.",
            "Không dùng tỷ lệ giữ chân của người nhận offer để kết luận offer có hiệu quả. Phải so sánh kết quả giữa những khách hàng có khả năng nhận can thiệp tương đương trong một thử nghiệm phù hợp.")
        self.detail_page("5 8 Thiết bị thuê bao và hồ sơ khách hàng", ["E18"], ["CurrentEquipmentDays", "HandsetPrice", "HandsetRefurbished", "ActiveSubs", "UniqueSubs", "CreditRating", "IncomeGroup", "AgeHH1"],
            "Tuổi thiết bị, mức giá thiết bị và số thuê bao phản ánh nhiều khía cạnh của quan hệ với công ty. Thiết bị cũ có thể phù hợp với nhu cầu hoặc là một tín hiệu cần hỗ trợ; số thuê bao hoạt động khác số thuê bao duy nhất cần được giải thích theo định nghĩa nguồn. Thu nhập và hồ sơ tín dụng có thể liên quan đến lựa chọn gói nhưng không tự tạo căn cứ để áp dụng một chính sách bất lợi.",
            "Bộ phận retention nên xác minh tình trạng thiết bị và lịch sử nâng cấp trước khi đề nghị đổi máy. Chính sách có trợ giá cần gắn với điều kiện kinh tế và nhu cầu thật. Các đặc điểm nhân khẩu học chỉ nên hỗ trợ hiểu phân khúc; khi triển khai phải rà soát sự khác biệt hiệu quả và cách phục vụ giữa các nhóm.",
            "Kết luận chương 5. Phân tích từng chủ đề tạo ra một tập tín hiệu có ý nghĩa nghiệp vụ. Các tín hiệu cần được kết hợp và kiểm tra ngoài mẫu trước khi dùng làm thứ tự liên hệ; mỗi hành động vẫn cần một phép đo tác động riêng.", compact=True)

    def chapter6(self):
        self.page("CHƯƠNG 6 PHÂN TÍCH ĐA BIẾN VÀ CẢNH BÁO", True)
        self.p("Một tín hiệu đơn lẻ có thể chỉ phản ánh đặc điểm của vòng đời hoặc mức sử dụng. Phân tích đa biến xem các tổ hợp để hiểu phạm vi của khác biệt, nhận diện các ô có ít khách hàng và tránh mô tả một mức churn cao như một quy tắc áp dụng cho mọi thuê bao.")
        first = self.chart(6, "E26", max_height=9.3, interpretation=False)
        second = next(c for c in self.charts if c.get('id') == 'E27')
        self.p(first['insight'])
        self.p(second['insight'])
        self.h("6 1 Hành trình rủi ro và điều kiện hành động")
        self.p("Các ô kết hợp cần được đọc cùng số khách hàng và mức nền chung. Một chênh lệch lớn ở ô rất nhỏ có thể không ổn định; một ô đông khách hàng với chênh lệch vừa phải có thể quan trọng hơn về nguồn lực. Phân tích này tìm phạm vi cần thử nghiệm, không xác nhận một tương tác nhân quả.")
        self.p("Hành trình chăm sóc bắt đầu bằng tín hiệu quan sát được, sau đó xác minh vấn đề rồi mới chọn hành động. Sự cố cần đối chiếu nhật ký kỹ thuật; thiết bị cần xác nhận tình trạng. Việc kết hợp này giúp chọn nội dung chẩn đoán, không xác định nguyên nhân churn.")
        self.p("Kết luận chương 6. Sự kết hợp các tín hiệu giúp đặt bối cảnh cho một trường hợp rủi ro. Mô hình ở Chương 8 sẽ kiểm tra liệu nhiều thông tin cùng nhau có cải thiện khả năng sắp hạng trên khách hàng chưa dùng để huấn luyện.")

    def chapter7(self):
        self.page("CHƯƠNG 7 XẾP HẠNG CÁC YẾU TỐ LIÊN QUAN", True)
        self.p("Xếp hạng liên hệ thống kê giúp thu hẹp tập yếu tố cần điều tra. Báo cáo tách mức liên hệ đơn biến khỏi tầm quan trọng dự báo đa biến; hai phép đo có thể khác nhau khi các trường cung cấp thông tin trùng lặp hoặc khi quan hệ phi tuyến.")
        self.chart(7, "E30")
        self.h("7 1 Cách đánh giá bằng chứng")
        self.p("Khi số quan sát lớn, một khác biệt nhỏ cũng có thể có ý nghĩa thống kê. Quyết định nghiệp vụ cần thêm độ lớn khác biệt, quy mô nhóm, độ ổn định và khả năng hành động. Với nhiều phép so sánh, cần đọc kết quả như bước khám phá; không chọn một chính sách chỉ vì một p value nhỏ.")
        self.p("Các trường liên quan đến retention được đánh dấu để tránh nhầm một chỉ báo của quy trình đã can thiệp với cảnh báo sớm. Các biến định danh bị loại khỏi mô hình; trường vùng dịch vụ cần được đánh giá theo ý nghĩa hoạt động và khả năng ghi nhớ cấu trúc mẫu. Mức liên hệ cao không tự chứng minh biến sẽ sử dụng ổn định ở tập khách hàng mới.")
        evidence = self.risk_summary.get('signal_evidence_train', {})
        if evidence:
            names = {'usage_decline':'sử dụng giảm', 'revenue_decline':'doanh thu giảm', 'equipment':'thiết bị dùng lâu', 'tenure_band':'nhóm vòng đời', 'quality':'chỉ báo chất lượng cao', 'care':'liên hệ chăm sóc nhiều'}
            accepted = [f"{names.get(k,k)} {percent(v['train_churn_rate'])}" for k,v in evidence.items() if v['included_as_risk_signal']]
            rejected = [f"{names.get(k,k)} {percent(v['train_churn_rate'])}" for k,v in evidence.items() if not v['included_as_risk_signal']]
            baseline = next(iter(evidence.values()))['train_baseline']
            self.p(f"Các cờ mô tả trên danh sách chấm điểm chỉ được giữ khi nhóm có đủ khách hàng và cận dưới Wilson 95% của churn rate trên train vượt mức nền {percent(baseline)}. Nhóm đáp ứng gồm " + '; '.join(accepted) + ". Nhóm không đáp ứng gồm " + '; '.join(rejected) + ". Đây là kiểm tra mô tả theo train, không phải giải thích cục bộ hoặc kiểm định tác động.")
        self.p("Kết luận chương 7. Chân dung rủi ro được xây từ nhiều bằng chứng thay vì một thuộc tính duy nhất. Danh sách yếu tố giải thích được dùng để chọn nội dung xác minh; mức đóng góp dự báo cần được kiểm tra bằng mô hình ngoài mẫu và được trình bày riêng ở chương sau.")

    def chapter8(self):
        m = self.model
        test = m["test"]
        self.page("CHƯƠNG 8 MÔ HÌNH DỰ ĐOÁN CHURN", True)
        self.p("Mô hình phục vụ cảnh báo trước khi khách hàng được đưa vào quy trình retention. Nhãn đầu ra là churn theo nguồn dữ liệu. Do không có đầy đủ mốc thời gian, kết quả là đánh giá ngoài mẫu trên ảnh chụp lịch sử; trước vận hành cần bổ sung đánh giá theo thời gian và xác định rõ cửa sổ quan sát, khoảng dự báo, ngày chấm điểm và thời điểm chốt kết quả.")
        self.table(8, "Phân chia dữ liệu có phân tầng", ["Phần dữ liệu", "Số khách hàng", "Số churn", "Tỷ lệ churn"], [[key, number(val["n"]), number(val["churn_count"]), percent(val["churn_rate"])] for key, val in m["splits"].items()], [4.8, 4, 4, 4])
        comparison = self.csv("model_comparison_validation.csv")
        cols = [c for c in ["model", "roc_auc", "pr_auc_average_precision", "precision_top10", "capture_top20", "brier_score"] if c in comparison]
        if cols:
            labels = {"model":"Mô hình", "roc_auc":"ROC AUC", "pr_auc_average_precision":"Average Precision", "precision_top10":"Precision top 10%", "capture_top20":"Capture top 20%", "brier_score":"Brier"}
            self.table(8, "So sánh mô hình trên validation", [labels[c] for c in cols], [[number(row[c], 3) if isinstance(row[c], (int, float)) else row[c] for c in cols] for _, row in comparison.iterrows()], [4.8] + [(16.8-4.8)/(len(cols)-1)]*(len(cols)-1), 8)
        self.h("8 1 Chọn mô hình và chuẩn bị đặc trưng")
        calibration = ("Hiệu chỉnh sigmoid được chọn sau kiểm tra trên validation." if m.get('calibrated') else "Không dùng hiệu chỉnh bên ngoài vì phép thử sigmoid không đáp ứng tiêu chí cải thiện trên validation.")
        self.p(f"Mô hình được chọn là {m['selected_model']}. {calibration} Brier trước hiệu chỉnh trên validation là {number(m['validation_brier_before'],4)}; phương án sigmoid cho {number(m['validation_brier_after_sigmoid'],4)}. Mô hình cơ sở và các mô hình cây được so sánh trên cùng cách chia dữ liệu. Các bước điền khuyết, mã hóa và chuẩn hóa nằm trong pipeline, bảo đảm phần kiểm tra không tham gia học tham số tiền xử lý.")
        self.p("Accuracy được trình bày cùng các chỉ số khác nhưng không quyết định mô hình một mình. ROC AUC đánh giá khả năng phân biệt thứ tự; Average Precision gắn với chất lượng precision và recall; Brier score đánh giá sai số xác suất; các chỉ số tại ngưỡng thể hiện hệ quả của một quyết định liên hệ cụ thể.")
        self.page("8 2 Hiệu quả trên tập kiểm tra độc lập")
        self.chart(8, "M02", max_height=6.8, interpretation=False)
        self.chart(8, "M05", width=13.0, max_height=6.6, interpretation=False)
        self.p(f"Trên test, ROC AUC là {number(test['roc_auc'],3)}, Average Precision là {number(test['pr_auc_average_precision'],3)}, accuracy là {percent(test['accuracy'])} và F1 là {number(test['f1'],3)}. Tại ngưỡng {number(test['threshold'],3)}, mô hình xác định đúng {number(test['tp'])} khách hàng churn, bỏ sót {number(test['fn'])}, đồng thời chọn {number(test['fp'])} khách hàng không churn.")
        self.p(f"Precision {percent(test['precision'])} cho biết phần khách hàng có nhãn churn trong nhóm mô hình dự báo dương; recall {percent(test['recall'])} cho biết phần churn được bao phủ. Những khách hàng dương tính giả làm tăng chi phí liên hệ, trong khi âm tính giả là phần rủi ro chưa được tiếp cận. Đây là đánh giá phân loại lịch sử, chưa phải số khách hàng sẽ được giữ lại bởi chiến dịch.")
        ci = m.get('test_bootstrap_95_ci', {}).get('roc_auc')
        if ci:
            self.p(f"Khoảng bootstrap 95% của ROC AUC trên test là {number(ci['lower_95'],3)} đến {number(ci['upper_95'],3)}. Khoảng này phản ánh biến thiên lấy mẫu trong tập lịch sử, chưa bao gồm thay đổi thị trường, thời gian hoặc cách đo lường.")
        self.page("8 3 Ngưỡng quyết định và giải thích mô hình")
        self.chart(8, "M07", max_height=6.9, interpretation=False)
        self.chart(8, "M04", width=15.8, max_height=6.9, interpretation=False)
        importance = self.csv('permutation_importance_test.csv')
        if not importance.empty:
            self.p("Các trường có mức giảm Average Precision lớn nhất khi hoán vị là " + '; '.join(f"{r['feature']} {number(r['importance_mean_pr_auc_decrease'],3)}" for _,r in importance.head(3).iterrows()) + ". Các giá trị này đo phụ thuộc dự báo trong mô hình hiện tại, không phải mức giảm churn có thể đạt bằng hành động.")
        self.p(f"Ngưỡng được chọn trên validation và được khóa trước khi đánh giá test. Kịch bản năng lực liên hệ dùng {percent(m.get('capacity_scenario', .2),0)} số khách hàng là giả định lập kế hoạch, không phải mức ngân sách hay năng lực đã được công ty xác nhận. Khi năng lực thay đổi, nên dùng thứ tự điểm và bảng targeting để lựa chọn quy mô phù hợp.")
        self.p("Permutation importance đo mức giảm chất lượng dự báo khi xáo trộn một trường trên dữ liệu kiểm tra. Nó không cho biết tác động nhân quả hoặc một hành động sửa trường đó sẽ giảm churn bao nhiêu. Các tín hiệu dễ trùng lặp có thể chia sẻ tầm quan trọng; tín hiệu rủi ro trên từng hồ sơ cần được dùng làm gợi ý xác minh thay vì lời giải thích chắc chắn cho quyết định của khách hàng.")
        self.p("Kết luận chương 8. Mô hình cung cấp một thứ tự rủi ro có thể đo lường ngoài mẫu. Lựa chọn vận hành cần ghép thứ tự đó với năng lực liên hệ, giá trị khách hàng và kết quả thử nghiệm giữ chân.")

    def chapter9(self):
        m = self.model
        self.page("CHƯƠNG 9 CHẤM ĐIỂM VÀ ƯU TIÊN KHÁCH HÀNG", True)
        self.p("Điểm rủi ro chuyển kết quả mô hình thành một thứ tự có thể sử dụng. Danh sách xuất kèm báo cáo là danh sách chấm điểm ảnh chụp lịch sử bằng dự đoán ngoài mẫu, không phải danh sách thuê bao hiện còn hoạt động để gọi ngay. Việc tạo danh sách chiến dịch cần lọc trạng thái hoạt động, điều kiện liên hệ, lịch sử tiếp xúc gần đây và thông tin có trước ngày can thiệp.")
        self.chart(9, "M11", "risk_distribution", "score_distribution", "risk_histogram", max_height=6.6, interpretation=False)
        bounds = m.get("risk_boundaries_validation_quantiles", {})
        if bounds:
            self.p("Các ngưỡng phân khúc được học từ phân phối điểm trên validation: " + "; ".join(f"{key} {number(val,3)}" for key,val in bounds.items()) + ". Tên phân khúc là quy ước ưu tiên nội bộ theo thứ tự điểm; không đồng nghĩa một khách hàng chắc chắn sẽ rời bỏ.")
        self.h("9 1 Giá trị khách hàng và quy tắc ưu tiên")
        self.p(f"MonthlyRevenue được dùng làm customer value proxy. Ngưỡng giá trị cao lấy từ trung vị tập huấn luyện là {number(m['high_value_threshold_train_median'],2)} theo đơn vị tiền nguồn. Dùng ngưỡng học trước giúp tránh nhìn vào kết quả test khi thiết kế phân khúc. Đại diện này chưa trừ chi phí mạng, hỗ trợ, khuyến mại hay rủi ro thanh toán nên không thể được gọi là lợi nhuận hoặc CLV.")
        self.p("Chính sách ưu tiên kết hợp rủi ro với giá trị, sau đó đối chiếu tín hiệu có thể xử lý. Một trường hợp điểm cao nhưng không còn thuê bao hoạt động phải được loại khỏi chiến dịch. Một khách hàng có dấu hiệu lỗi dịch vụ cần được chuyển xử lý phù hợp trước khi cân nhắc giá, còn khách hàng giá trị thấp nên được tiếp cận qua kênh có chi phí tương xứng.")
        if self.candidate_stats:
            c = self.candidate_stats
            self.p(f"Trong danh sách lịch sử, {number(c['all_p1'])} hồ sơ thuộc P1 nhưng chỉ {number(c['candidate_p1'])} hồ sơ đồng thời chưa churn theo nhãn quan sát. Cột điều kiện liên hệ lịch sử đánh dấu việc sàng lọc này; trạng thái hoạt động hiện tại, lịch sử liên hệ và điều kiện nhận ưu đãi cần được xác minh lại trước khi phân công.")
        self.page("9 2 Ma trận rủi ro và giá trị")
        self.chart(9, "M10", max_height=7.8, interpretation=False)
        risk = ROOT / "Cell2Cell_Churn_Risk_Scores.csv"
        if not risk.exists():
            risk = ROOT / "output" / risk.name
        if risk.exists():
            scores = pd.read_csv(risk)
            if "risk_segment" in scores:
                rows = []
                for group, block in scores.groupby("risk_segment", sort=False):
                    rows.append([group, number(len(block)), percent(block["churn_probability"].mean()), number(block["customer_value_proxy"].median(),2) if "customer_value_proxy" in block else ""])
                self.table(9, "Quy mô phân khúc trong danh sách lịch sử", ["Phân khúc", "Khách hàng", "Điểm trung bình", "Trung vị giá trị"], rows, [5,4,4,3.8])
        personas = self.risk_summary.get('personas', [])
        if personas:
            self.p("Các chân dung được gán theo thứ tự quy tắc nên không chồng lặp. " + '; '.join(f"{p['customer_persona']} gồm {number(p['customers'])} khách hàng, điểm trung bình {percent(p['mean_probability'])}" for p in personas) + ". Nhãn chân dung là mô tả lịch sử nhằm chọn nội dung liên hệ, không phải kết quả kiểm chứng một offer.")
        self.p("Nhóm rủi ro cao và giá trị cao cần được rà soát trước vì kết hợp khả năng rời bỏ với nguồn doanh thu đáng quan tâm. Nhóm rủi ro cao và giá trị thấp vẫn cần xử lý vấn đề trải nghiệm nhưng lựa chọn kênh và mức ưu đãi phải tôn trọng chi phí. Nhóm rủi ro thấp và giá trị cao phù hợp với chăm sóc quan hệ, không cần một chiến dịch giảm giá phản ứng.")
        self.p("Kết luận chương 9. Ma trận hỗ trợ phân công và thiết kế chính sách; nó không tự tối ưu lợi nhuận. Bước vận hành kế tiếp là đo độ bao phủ tại từng mức năng lực, sau đó thử nghiệm giá trị tăng thêm của hành động.")

    def chapter10(self):
        self.page("CHƯƠNG 10 HIỆU QUẢ LỰA CHỌN KHÁCH HÀNG", True)
        self.p("Khả năng sắp hạng có ý nghĩa khi công ty không thể liên hệ mọi khách hàng. Phân tích targeting xếp khách hàng test theo điểm giảm dần và tính chất lượng tại các mức bao phủ. Các kết quả này không dùng nhãn test để thay đổi mô hình hoặc chọn lại ngưỡng.")
        self.chart(10, "M09", interpretation=False)
        table = self.csv("targeting_test.csv")
        if not table.empty:
            rows = [[percent(r['capacity'],0),number(r['customers']),percent(r['precision']),percent(r['capture_recall']),number(r['lift'],2)] for _,r in table.iterrows()]
            self.table(10, "Hiệu quả tại các mức năng lực liên hệ trên test", ["Bao phủ", "Khách hàng", "Precision", "Churn được phủ", "Lift"], rows, [3,3.5,3.5,3.8,3])
        self.p("Precision trong nhóm được chọn cho biết mức tập trung churn lịch sử; cumulative gains cho biết phần churn được bao phủ khi mở rộng liên hệ. Lift so sánh precision đó với tỷ lệ churn nền. Cả ba đo chất lượng lựa chọn, không đo khả năng khách hàng thay đổi quyết định sau một offer.")
        self.p("Bảng trên chọn đúng số khách hàng theo tỷ lệ bao phủ và thứ tự điểm. Kết quả có thể lệch nhẹ so với ma trận nhầm lẫn ở Chương 8 vì ma trận áp một ngưỡng cố định học từ validation, trong khi bảng targeting lấy đúng thứ hạng trên test. Hai cách trả lời hai quyết định khác nhau và không được đổi mẫu số để làm khớp kết quả.")
        self.p("Kết luận chương 10. Bộ phận retention có thể dùng bảng bao phủ để lập phương án nguồn lực. Quy mô cuối cùng phải phụ thuộc chi phí, năng lực xử lý và giá trị giữ chân tăng thêm được đo bằng thử nghiệm, thay vì lấy một tỷ lệ liên hệ cố định làm mục tiêu lâu dài.")

    def chapter11(self):
        self.page("CHƯƠNG 11 CHÍNH SÁCH GIỮ CHÂN KHÁCH HÀNG", True)
        self.p("Chính sách được đề xuất như các giả thuyết cần thử nghiệm, dựa trên tín hiệu quan sát và vấn đề có thể xử lý. Phân khúc xác định đối tượng kiểm tra; cuộc liên hệ xác minh nhu cầu; hành động phải tương xứng với vấn đề, giá trị và chi phí. Báo cáo không gán một khoản ngân sách hoặc tỷ lệ thành công chưa có trong dữ liệu.")
        comps = self.summary.get('core_comparisons', {})
        contrasts = []
        for key in ['CustomerCareCalls', 'DroppedCalls']:
            if key in comps:
                churn, other = comps[key]['1']['mean'], comps[key]['0']['mean']
                contrasts.append(f"{key} trung bình nhóm churn {number(churn,2)}, {'thấp hơn' if churn < other else 'cao hơn' if churn > other else 'bằng'} nhóm duy trì {number(other,2)}")
        if contrasts:
            self.p("Bằng chứng cần lưu ý là " + '; '.join(contrasts) + ". Vì vậy không mặc định khách hàng gọi chăm sóc nhiều hoặc có nhiều cuộc gọi rớt là nhóm cần ưu tiên churn. Xử lý vấn đề trải nghiệm vẫn cần thiết khi được xác minh, còn gắn tín hiệu rủi ro phải qua kiểm tra trên tập huấn luyện.")
        self.table(11, "Chính sách theo tín hiệu và tình trạng khách hàng", ["Nhóm", "Nội dung xác minh", "Hành động đề xuất", "KPI chính"], [
            ["Giá trị cao và điểm rủi ro cao", "Nhu cầu, gói cước và trải nghiệm", "Rà soát cá nhân và giải quyết vấn đề đã xác minh", "Giữ chân tăng thêm so với đối chứng"],
            ["Sử dụng giảm", "Thay đổi nhu cầu hoặc trải nghiệm", "Liên hệ chẩn đoán và đề nghị gói phù hợp", "Sử dụng phục hồi và retention lift"],
            ["Doanh thu giảm cần xác minh", "Lịch sử gói, cước và nhu cầu", "Rà soát gói; chưa dùng làm tín hiệu rủi ro chính", "Doanh thu đóng góp tăng thêm"],
            ["Tín hiệu chất lượng", "Lỗi mạng hoặc vùng phủ", "Xử lý kỹ thuật và phản hồi tiến độ", "Sự cố lặp lại và retention lift"],
            ["Yêu cầu hỗ trợ chưa giải quyết", "Nội dung và kết quả xử lý", "Một đầu mối xử lý và xác nhận kết quả", "Tái liên hệ và thời gian giải quyết"],
            ["Thiết bị cần rà soát", "Lỗi máy và nhu cầu nâng cấp", "Hỗ trợ kỹ thuật hoặc nâng cấp có điều kiện", "Chi phí trên giữ chân tăng thêm"],
        ], [3.5,4.5,5.7,3.1], 8.5)
        self.h("11 1 Khách hàng rủi ro cao và giá trị cao")
        self.p("Nhóm này cần một đầu mối có đủ thẩm quyền xử lý vấn đề và kiểm tra lịch sử trước khi liên hệ. Tín hiệu quan sát quyết định nội dung chẩn đoán; giá trị quyết định giới hạn chi phí có thể xem xét. Đội retention nên ghi lại lý do liên hệ, nhu cầu xác nhận, hành động thực tế và kết quả để chuyển từ dự báo churn sang hiểu phản ứng với chính sách.")
        self.p("Đối với nhóm có điểm thấp hoặc không có vấn đề cụ thể, chăm sóc quan hệ và thông tin gói có thể phù hợp hơn một ưu đãi giá. Việc gửi ưu đãi cho những người vốn vẫn duy trì dịch vụ tạo chi phí mà không tăng giữ chân. Nhóm đối chứng giúp phát hiện phần chi phí này.")
        self.page("11 2 Ưu tiên hành động và nguyên tắc kinh tế")
        self.table(11, "Ma trận ưu tiên triển khai", ["Ưu tiên", "Điều kiện theo danh sách điểm", "Hành động", "KPI"], [
            ["P1", "Giá trị cao và rủi ro cao hoặc rất cao", "Rà soát cá nhân; xử lý vấn đề đã xác minh", "Giữ chân tăng thêm và chi phí"],
            ["P2", "Giá trị cao và rủi ro trung bình", "Chăm sóc theo nhu cầu; xác minh tín hiệu", "Retention lift và giá trị tăng thêm"],
            ["P3", "Giá trị thấp hơn ngưỡng, rủi ro trung bình hoặc cao", "Kênh chi phí thấp và offer có chọn lọc", "Chi phí trên giữ chân tăng thêm"],
            ["P4", "Các hồ sơ còn lại", "Theo dõi và duy trì quan hệ", "Churn rate và chi phí chăm sóc"],
        ], [2.5,5.2,5,4.1])
        self.p("Giá trị kỳ vọng của một liên hệ cần dựa trên xác suất giữ chân tăng thêm do hành động, phần đóng góp kinh tế nếu giữ được khách hàng, chi phí liên hệ và chi phí offer. Xác suất churn dự báo chưa phải xác suất tác động của offer. Chỉ nên dùng điểm churn để sắp hạng trước khi có dữ liệu thử nghiệm đủ tốt để học tác động khác biệt giữa các khách hàng.")
        self.p("Công thức quyết định là lợi ích tăng thêm kỳ vọng bằng mức tăng xác suất duy trì do can thiệp nhân với giá trị đóng góp trong kỳ đánh giá, trừ chi phí liên hệ và chi phí ưu đãi kỳ vọng. ROI bằng lợi ích ròng tăng thêm chia cho tổng chi phí chiến dịch. Các thành phần chưa có trong nguồn dữ liệu cần được đo hoặc cung cấp trước khi tính kết quả bằng tiền.")
        self.p("Kết luận chương 11. Ưu tiên trước là xử lý đúng vấn đề cho khách hàng đáng quan tâm, sau đó mới tối ưu mức offer. Một chính sách chỉ nên mở rộng khi chứng minh được giữ chân tăng thêm và hiệu quả kinh tế so với nhóm đối chứng.")

    def chapter12(self):
        self.page("CHƯƠNG 12 TRIỂN KHAI VÀ ĐO LƯỜNG", True)
        self.p("Quy trình vận hành cần tách dữ liệu phục vụ dự báo khỏi dữ liệu tạo ra sau can thiệp. Mỗi đợt chấm điểm phải có phiên bản dữ liệu, thời điểm trích xuất, phiên bản mô hình và tiêu chí đủ điều kiện liên hệ để kết quả có thể kiểm tra lại.")
        self.table(12, "Quy trình vận hành đề xuất", ["Bước", "Nội dung", "Dấu vết cần lưu"], [
            ["Cập nhật dữ liệu", "Chốt hồ sơ hoạt động và cửa sổ quan sát", "Ngày chốt, số hồ sơ, thiếu dữ liệu"],
            ["Chấm điểm", "Áp pipeline và mô hình đã khóa", "Phiên bản, điểm, tín hiệu, dữ liệu đầu vào"],
            ["Ưu tiên", "Kết hợp rủi ro, giá trị, năng lực và điều kiện", "Lý do chọn, phân khúc, loại trừ"],
            ["Phân nhóm thử nghiệm", "Ngẫu nhiên trong từng tầng rủi ro và giá trị", "Mã nhóm và xác suất phân nhóm"],
            ["Can thiệp", "Xử lý vấn đề hoặc offer theo chính sách", "Tiếp xúc, hành động, chi phí, chấp nhận"],
            ["Theo dõi", "Chốt churn và giá trị sau can thiệp", "Ngày, trạng thái, doanh thu đóng góp"],
            ["Đánh giá và cập nhật", "So sánh đối chứng, drift và chất lượng điểm", "Báo cáo hiệu quả và quyết định thay đổi"],
        ], [3.2,7.0,6.6])
        self.p("Người phụ trách dữ liệu cần theo dõi thay đổi tỷ lệ thiếu, phân phối đặc trưng và cơ cấu khách hàng. Nhóm mô hình kiểm tra chất lượng sắp hạng và hiệu chỉnh sau khi nhãn mới hoàn tất; nhóm retention theo dõi tiếp xúc, xử lý và chi phí. Một mô hình có điểm số ổn định nhưng chiến dịch không tạo lift cần được xem xét ở nội dung hành động và cách thực thi.")
        self.h("12 1 Thử nghiệm có đối chứng")
        self.p("Khách hàng đủ điều kiện nên được ngẫu nhiên vào treatment và control trong từng tầng rủi ro và giá trị. Hai nhóm được theo dõi trên cùng khoảng thời gian và cùng quy tắc xác định churn. Phân tích chính theo nhóm được phân ban đầu để tránh chỉ đánh giá những người đã trả lời hoặc đã nhận offer. Cỡ mẫu cần được tính từ tỷ lệ nền, mức cải thiện tối thiểu cần phát hiện và công suất mong muốn trước khi chạy.")
        self.page("12 2 Hệ thống KPI và tiêu chuẩn mở rộng")
        self.table(12, "KPI gắn với dự báo và tác động", ["KPI", "Công thức", "Ý nghĩa và tần suất"], [
            ["Churn rate", "Số churn chia số đủ điều kiện trong kỳ", "Theo tổng thể, rủi ro và giá trị mỗi kỳ nhãn"],
            ["Precision nhóm ưu tiên", "Churn trong nhóm được chọn chia số được chọn", "Chất lượng xếp hạng mỗi đợt nhãn hoàn tất"],
            ["High risk capture", "Churn trong nhóm được chọn chia tổng churn", "Độ bao phủ theo năng lực liên hệ"],
            ["Lift", "Precision nhóm chọn chia churn rate nền", "So với lựa chọn ngẫu nhiên"],
            ["Contact rate", "Khách hàng liên hệ thành công chia được phân công", "Khả năng thực thi hằng chiến dịch"],
            ["Offer acceptance", "Khách hàng nhận offer chia khách hàng được đề nghị", "Phản ứng với offer theo chính sách"],
            ["Incremental retention", "Tỷ lệ duy trì treatment trừ control", "Tác động tăng thêm khi hết kỳ theo dõi"],
            ["Chi phí giữ chân tăng thêm", "Tổng chi phí chia số duy trì tăng thêm", "Chỉ tính khi hiệu quả tăng thêm dương"],
            ["Giá trị được bảo vệ", "Đóng góp treatment trừ control sau quy đổi", "Cùng định nghĩa chi phí và kỳ đánh giá"],
            ["Duy trì theo cửa sổ", "Trạng thái sau các mốc theo dõi đã định", "Đo ngắn và dài hạn theo thiết kế thử nghiệm"],
        ], [4,6.0,6.8], 8.5)
        self.p("Incremental retention lift bằng tỷ lệ duy trì của treatment trừ tỷ lệ duy trì của control; tương đương churn rate control trừ churn rate treatment khi dùng cùng mẫu số. So sánh trước và sau một chiến dịch không kiểm soát được thay đổi nhu cầu, mùa vụ và cách chọn khách hàng, nên không đủ để xác nhận tác động.")
        self.p("Kết luận chương 12. Mở rộng chương trình cần đồng thời đáp ứng giữ chân tăng thêm, hiệu quả chi phí và chất lượng thực thi. Các mốc theo dõi ngắn và dài hạn là tham số thiết kế chính sách cần được thống nhất trước; báo cáo không suy ra chúng từ dữ liệu ảnh chụp.")

    def chapter13(self):
        s, m = self.summary, self.model
        self.page("CHƯƠNG 13 KẾT LUẬN VÀ KIẾN NGHỊ", True)
        self.p(f"Dữ liệu khách hàng của công ty ghi nhận tỷ lệ churn {percent(s['churn_rate'])} trên {number(s['n_customers'])} hồ sơ có nhãn. Phân tích đã kết nối quy mô churn, các nhóm hành vi và một mô hình cảnh báo sớm; trên test, mô hình {m['selected_model']} đạt ROC AUC {number(m['test']['roc_auc'],3)} và Average Precision {number(m['test']['pr_auc_average_precision'],3)}. Các kết quả hỗ trợ lựa chọn khách hàng cần xem xét trước, chưa xác nhận tác động của một chính sách giữ chân.")
        comps = s['core_comparisons']
        tenure_rows = [r for r in self.get_analysis('MonthsInService')['table'] if r['n'] >= 100]
        tenure = max(tenure_rows, key=lambda r:r['churn_rate'])
        missing_change = next(r for r in self.get_analysis('PercChangeMinutes')['table'] if 'thiếu' in str(r['group']).lower())
        c = self.candidate_stats
        self.p(f"Năm phát hiện cần giữ lại là quy mô churn {percent(s['churn_rate'])} trong cohort; thiết bị của nhóm churn có trung vị {number(comps['CurrentEquipmentDays']['1']['median'])} ngày so với {number(comps['CurrentEquipmentDays']['0']['median'])}; nhóm thiếu chỉ số thay đổi sử dụng có churn {percent(missing_change['churn_rate'])} và cần kiểm tra cơ chế đo; nhóm vòng đời {tenure['group']} có churn {percent(tenure['churn_rate'])}; và mô hình sắp hạng tạo được một nhóm ứng viên P1 lịch sử {number(c['candidate_p1'])} khách hàng chưa churn cần xác minh lại trạng thái.")
        rows = [
            ['Thiết bị', f"Trung vị churn {number(comps['CurrentEquipmentDays']['1']['median'])} so với {number(comps['CurrentEquipmentDays']['0']['median'])} ngày", 'Thiết bị dùng lâu', 'P1–P3 theo điểm', 'Kiểm tra tình trạng và nhu cầu đổi thiết bị', 'Retention lift; chi phí'],
            ['Thay đổi sử dụng', f"Thiếu {number(missing_change['n'])} hồ sơ; churn {percent(missing_change['churn_rate'])}", 'Thiếu chỉ số hoặc sử dụng giảm', 'Dữ liệu và P1–P3', 'Kiểm tra đo lường; xác minh nhu cầu ở nhóm sử dụng giảm', 'Missing rate; retention lift'],
            ['Vòng đời', f"Nhóm {tenure['group']} có churn {percent(tenure['churn_rate'])}", f"Gắn bó {tenure['group']} tháng", 'P1–P3 theo điểm', 'Rà soát nhu cầu và trải nghiệm theo giai đoạn quan hệ', 'Churn theo vòng đời'],
            ['Mức sử dụng', f"Trung vị churn {number(comps['MonthlyMinutes']['1']['median'])} so với {number(comps['MonthlyMinutes']['0']['median'])} phút", 'Mức dùng thấp cần xác minh', 'P1–P3 theo điểm', 'Kiểm tra nhu cầu và tính phù hợp của gói', 'Sử dụng; retention lift'],
            ['Doanh thu', f"Trung vị churn {number(comps['MonthlyRevenue']['1']['median'],2)} so với {number(comps['MonthlyRevenue']['0']['median'],2)}", 'Phân khúc theo giá trị', 'Trước thử nghiệm', 'Rà soát gói và biên lợi nhuận trước ưu đãi', 'Giá trị tăng thêm'],
            ['Sắp hạng rủi ro', f"Test ROC AUC {number(m['test']['roc_auc'],3)}; AP {number(m['test']['pr_auc_average_precision'],3)}", f"{number(c['candidate_p1'])} ứng viên P1 lịch sử", 'Thử nghiệm P1', 'Xác minh đang hoạt động; chạy pilot có đối chứng', 'Precision; retention lift'],
        ]
        self.table(13, "Hành động điều hành gắn với bằng chứng", ["Phát hiện", "Bằng chứng", "Nhóm", "Ưu tiên", "Hành động", "KPI"], rows, [2.2,4.0,2.7,1.9,3.7,2.3], 7.8)
        self.h("13 1 Các hành động nên thực hiện trước")
        self.p("Thứ nhất, chốt định nghĩa churn, cửa sổ dự báo và tập thuê bao đủ điều kiện liên hệ. Thứ hai, rà soát các biến theo thời điểm để bảo đảm cảnh báo sớm không dùng kết quả retention. Thứ ba, tạo nhóm ưu tiên từ điểm rủi ro và giá trị, đồng thời phân luồng theo vấn đề cần xác minh. Thứ tư, thử nghiệm một số chính sách có đối chứng trong từng tầng rủi ro và giá trị. Thứ năm, ghi đầy đủ chi phí và kết quả để quyết định chính sách nào được mở rộng.")
        self.page("13 2 Hạn chế và hướng phát triển")
        self.p("Dữ liệu quan sát chỉ phản ánh các trường đã được ghi nhận. Một khác biệt churn có thể chịu tác động của nhu cầu, chất lượng vùng phủ, lịch sử cước, đối thủ hoặc trải nghiệm không xuất hiện trong dữ liệu. Việc gọi một yếu tố là tín hiệu rủi ro không đồng nghĩa yếu tố đó gây ra churn hoặc sửa yếu tố đó sẽ thay đổi quyết định của khách hàng.")
        self.p("Ảnh chụp dữ liệu thiếu lịch sử dọc đầy đủ để kiểm tra thứ tự thời gian và cửa sổ dự báo. Phân tách ngẫu nhiên có phân tầng giúp đánh giá ngoài mẫu trong cùng nguồn lịch sử, nhưng chưa đo được độ ổn định khi mạng, sản phẩm hoặc cơ cấu khách hàng thay đổi. Cần bổ sung một tập kiểm tra theo thời gian trước khi dùng điểm cho hoạt động thường xuyên.")
        self.p("MonthlyRevenue là đại diện giá trị sẵn có, chưa chứa chi phí phục vụ hay giá trị vòng đời. Tỷ lệ churn, phần doanh thu gắn với churn và precision của mô hình không đủ để suy ra doanh thu có thể được bảo vệ. Thông tin offer, biên lợi nhuận, chi phí liên hệ và tác động thử nghiệm là điều kiện cần cho tính toán kinh tế.")
        self.h("13 3 Dữ liệu cần bổ sung và bước chuyển tiếp")
        self.p("Ưu tiên bổ sung mốc kích hoạt và churn, lịch sử hóa đơn, trạng thái thuê bao, nhật ký chất lượng mạng theo vùng và thời gian, ticket hỗ trợ và nội dung khiếu nại, phản hồi trải nghiệm, hoạt động trên ứng dụng, thay đổi gói và kết quả chiến dịch. Các nguồn cần được nối theo khách hàng và thời điểm với quy tắc kiểm soát tính sẵn có trước can thiệp.")
        self.p("Khi có dữ liệu thử nghiệm, công ty có thể đi tiếp từ dự báo ai dễ churn sang dự báo ai thay đổi hành vi khi nhận từng hành động. Mô hình tác động khác biệt và tối ưu chi phí khi đó mới có cơ sở để thay thế quy tắc ưu tiên dựa trên rủi ro và giá trị. Việc cập nhật cần duy trì một phép đo đối chứng, kiểm tra chất lượng phục vụ giữa các nhóm và lưu phiên bản để truy vết.")
        self.p("Kết luận chương 13. Công ty nên bắt đầu bằng một chương trình có giới hạn, hướng đến khách hàng vừa có rủi ro vừa có giá trị và có vấn đề cụ thể để xử lý. Thành công được xác định bằng số khách hàng được giữ lại thêm và hiệu quả kinh tế so với đối chứng; điểm dự báo là công cụ chọn đối tượng cho quá trình đó.")
        self.h('Tài liệu tham khảo', 2)
        self.reference('Knaflic C N (2015). Storytelling with Data A Data Visualization Guide for Business Professionals. Wiley.', 'Bản tài liệu do người dùng cung cấp', 'https://luutruvn.com/wp-content/uploads/2025/03/Storytelling-with-Data_-A-Data-Visualization-Guide-for-Business-Professionals-PDFDrive-.pdf')
        self.reference('Ward M O, Grinstein G và Keim D (2015). Interactive Data Visualization Foundations Techniques and Applications. Ấn bản thứ hai.', 'Trang của tác giả', 'https://www.idvbook.com/index.html')
        self.reference('Empulse. Tài liệu bộ dữ liệu Cell2Cell. Truy cập ngày 27 tháng 10 năm 2026.', 'Tài liệu dữ liệu', 'https://empulse.readthedocs.io/en/latest/guide/datasets/cell2cell.html')
        self.reference('Empulse. Mã nguồn bộ tải và quy tắc xử lý Cell2Cell. Truy cập ngày 27 tháng 10 năm 2026.', 'Mã nguồn bộ tải', 'https://empulse.readthedocs.io/en/latest/_modules/empulse/datasets/_remote.html')
        self.reference('Cell2Cell. Tệp nguồn được ghim theo phiên bản và đối chiếu SHA256 trong data manifest.', 'Dữ liệu nguồn', self.data_manifest['source_url'])

    def build(self):
        required_summary = ["n_customers", "n_features", "n_numeric", "n_categorical", "n_churn", "churn_rate"]
        missing = [x for x in required_summary if x not in self.summary]
        if missing or "test" not in self.model:
            raise ValueError(f"Analysis inputs incomplete: summary missing {missing}; model test {'test' in self.model}")
        self.intro_pages()
        self.executive()
        for number_ in range(1,14):
            getattr(self, f"chapter{number_}")()
        self.output.parent.mkdir(parents=True,exist_ok=True)
        self.doc.save(self.output)
        manifest = {"docx": str(self.output), "figure_count": len(self.figures), "table_count": len(self.tables), "figures": self.figures, "tables": self.tables, "chart_ids": sorted(str(x) for x in self.used_charts), "intended_pages": 34}
        (self.analysis / "report_manifest.json").write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding="utf-8")
        print(json.dumps(manifest,ensure_ascii=False,indent=2))


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser()
    parser.add_argument("--analysis-dir",type=Path,default=ANALYSIS)
    parser.add_argument("--output",type=Path,default=OUT)
    args = parser.parse_args()
    Report(args.analysis_dir,args.output).build()
